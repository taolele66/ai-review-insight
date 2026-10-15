"""第 2–3 步：抽样，生成两份 Excel。

1. pilot_50.xlsx：读 50 条用来修订标签体系（不进评估集）
2. annotation_300.xlsx：评估集，按星级分层抽样，每条留好下拉框，人工填写后作为「标准答案」

分层：Amazon 评分高度偏 5 星，纯随机抽 300 条差评只有几十条，痛点几乎没法评估。
这里按星级等量抽（每档 60 条），报告整体准确率时按真实星级分布加权（evaluate.py 里已处理）。

用法：python -m src.eval.make_annotation_sheet
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from src import config as C
from src.labeling.schema import ASPECTS, PAIN_POINTS, PURCHASE_MOTIVES, REPURCHASE_INTENT, SKIN_TYPES

PER_STAR = 60


def sample(df):
    pilot = df.sample(50, random_state=C.SEED + 1)
    rest = df.drop(pilot.index)
    evalset = pd.concat([g.sample(min(PER_STAR, len(g)), random_state=C.SEED)
                         for _, g in rest.groupby("rating")])
    evalset = evalset.sample(frac=1, random_state=C.SEED)  # 打乱，标注时看不出星级顺序
    return pilot, evalset


def write_sheet(df, path, with_labels):
    wb = Workbook()
    ws = wb.active
    ws.title = "标注"
    aspect_cols = [f"aspect_{a}" for a in ASPECTS]
    head = ["review_id", "title", "text"]
    if with_labels:
        head += aspect_cols + ["pain_points", "purchase_motive", "skin_type", "repurchase_intent", "is_relevant",
                              "overall_sentiment", "hard_case", "note"]
    else:
        head += ["我的观察（这条在谈什么？现有标签够不够用？）"]
    ws.append(head)
    for r in df.itertuples():
        ws.append([r.review_id, r.title, r.text])

    bold = Font(bold=True)
    for c in ws[1]:
        c.font = bold
        c.fill = PatternFill("solid", fgColor="DDE6F0")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 70
    for row in ws.iter_rows(min_row=2, max_col=3):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"

    if with_labels:
        n = len(df) + 1
        sent = DataValidation(type="list", formula1='"pos,neu,neg"', allow_blank=True)
        motive = DataValidation(type="list", formula1='"' + ",".join(PURCHASE_MOTIVES) + '"', allow_blank=True)
        overall = DataValidation(type="list", formula1='"pos,neu,neg"', allow_blank=True)
        hard = DataValidation(type="list", formula1='"反讽,多方面混杂,很短,其他"', allow_blank=True)
        intent = DataValidation(type="list", formula1='"yes,no,none"', allow_blank=True)
        relevant = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
        for dv in (sent, motive, overall, hard, intent, relevant):
            ws.add_data_validation(dv)
        for j, col in enumerate(head, 1):
            letter = ws.cell(1, j).column_letter
            if col.startswith("aspect_"):
                ws.cell(1, j).value = f"{col}\n{ASPECTS[col[7:]][0]}"
                ws.column_dimensions[letter].width = 11
                sent.add(f"{letter}2:{letter}{n}")
            elif col == "purchase_motive":
                motive.add(f"{letter}2:{letter}{n}"); ws.column_dimensions[letter].width = 16
            elif col == "overall_sentiment":
                overall.add(f"{letter}2:{letter}{n}"); ws.column_dimensions[letter].width = 12
            elif col == "hard_case":
                hard.add(f"{letter}2:{letter}{n}")
            elif col == "repurchase_intent":
                intent.add(f"{letter}2:{letter}{n}"); ws.column_dimensions[letter].width = 12
            elif col == "is_relevant":
                relevant.add(f"{letter}2:{letter}{n}"); ws.column_dimensions[letter].width = 10
            elif col in ("pain_points", "skin_type"):
                ws.column_dimensions[letter].width = 22
        ws.row_dimensions[1].height = 45

    guide = wb.create_sheet("标签说明")
    guide.append(["类型", "代码", "中文", "说明"])
    for k, (zh, desc) in ASPECTS.items():
        guide.append(["方面 aspect", k, zh, desc])
    for k, zh in PAIN_POINTS.items():
        guide.append(["痛点 pain_points（多个用英文逗号分隔）", k, zh, ""])
    for k, zh in PURCHASE_MOTIVES.items():
        guide.append(["购买动机 purchase_motive（单选）", k, zh, ""])
    for k, zh in SKIN_TYPES.items():
        guide.append(["肤质 skin_type（多个用英文逗号分隔）", k, zh, ""])
    for k, zh in REPURCHASE_INTENT.items():
        guide.append(["复购意向 repurchase_intent（单选）", k, zh, ""])
    guide.append(["是否相关 is_relevant", "TRUE / FALSE", "", "看商品本身是不是护肤品：一次性口罩、汽车配件、头发产品等填 FALSE；"
                  "护肤品一律 TRUE（即使评论只说物流、数量、包装）。留空视为 TRUE"])
    guide.append(["复购意向规则", "", "", "只有明确表达才填：yes = 会再买/买过多次/一直在用；no = 明确不会再买、已退货、换掉了；"
                  "只是差评但没说买不买 = none"])
    guide.append([])
    guide.append(["填写规则", "", "", "方面列：只在评论谈到该方面时填 pos/neu/neg，没谈到留空。"
                  "pain_points / skin_type：填代码，多个用逗号分隔，没有留空。repurchase_intent 留空视为 none。hard_case：遇到反讽、多方面混杂、很短等难例时标一下，"
                  "之后用来分析 AI 在哪类评论上容易错。"])
    for col, w in zip("ABCD", (34, 20, 18, 70)):
        guide.column_dimensions[col].width = w
    wb.save(path)


def main():
    df = pd.read_parquet(C.PROC_DIR / "reviews_clean.parquet")
    pilot, evalset = sample(df)
    ann = C.ROOT / "annotation"
    ann.mkdir(exist_ok=True)
    write_sheet(pilot, ann / "pilot_50.xlsx", with_labels=False)
    write_sheet(evalset, ann / "annotation_300.xlsx", with_labels=True)
    (C.PROC_DIR / "eval_ids.txt").write_text("\n".join(evalset["review_id"]))
    (C.PROC_DIR / "pilot_ids.txt").write_text("\n".join(pilot["review_id"]))
    print(f"pilot {len(pilot)} 条，评估集 {len(evalset)} 条，星级分布：")
    print(evalset["rating"].value_counts().sort_index().to_string())


if __name__ == "__main__":
    main()
